# /// script
# dependencies = ["openpyxl"]
# ///
"""Build the canonical "ad hoc sheet" for the Scratchwork demo video.

Run: uv run docs/demo-assets/make_scratch_xlsx.py
Same six months as the tutorials, so every answer ties to the tour workbook.
The SUMIF range deliberately stops one row short: June was added after the
formula was written. That is the bug the video is about.
"""
from openpyxl import Workbook
from openpyxl.styles import Font

wb = Workbook()
ws = wb.active
ws.title = "Scratch"
bold = Font(bold=True)

ws["A1"], ws["B1"] = "Growth", 1.08
ws["A2"], ws["B2"] = "Tax rate", 0.26

ws["A4"], ws["B4"], ws["C4"] = "Month", "Revenue", "Region"
for c in ("A4", "B4", "C4"):
    ws[c].font = bold
rows = [
    ("2026-01", 118000, "East"),
    ("2026-02", 124000, "West"),
    ("2026-03", 136000, "East"),
    ("2026-04", 142000, "West"),
    ("2026-05", 151000, "East"),
    ("2026-06", 168000, "East"),  # added later; the SUMIF below never saw it
]
for i, (m, r, g) in enumerate(rows, start=5):
    ws[f"A{i}"], ws[f"B{i}"], ws[f"C{i}"] = m, r, g

ws["E4"], ws["F4"] = "Total", "=SUM(B5:B10)"
ws["E5"], ws["F5"] = "East total", '=SUMIF(C5:C9,"East",B5:B9)'
ws["E6"], ws["F6"] = "East share", "=F5/F4"
ws["E7"], ws["F7"] = "Forecast", "=ROUND(F4*B1,0)"
ws["E8"], ws["F8"] = "After tax", "=F7*(1-B2)"
for c in ("E4", "E5", "E6", "E7", "E8"):
    ws[c].font = bold
for c in ("B5", "B6", "B7", "B8", "B9", "B10", "F4", "F5", "F7", "F8"):
    ws[c].number_format = "#,##0"
ws["F6"].number_format = "0.0%"
ws["B2"].number_format = "0%"
ws.column_dimensions["A"].width = 11
ws.column_dimensions["E"].width = 11

wb.save("docs/demo-assets/scratch.xlsx")
print("wrote docs/demo-assets/scratch.xlsx")
